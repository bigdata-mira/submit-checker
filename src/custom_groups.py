from __future__ import annotations

from collections.abc import Callable, Iterable, Sequence

from openpyxl import load_workbook
from werkzeug.datastructures import FileStorage

from src.target_manager import CustomGroup, FolderScanError, get_custom_group_names, load_target_store, save_custom_groups

HEADER_HINTS: frozenset[str] = frozenset(
    {
        "대상",
        "대상명",
        "취합대상",
        "취합대상명",
        "이름",
        "성명",
        "명단",
        "소속",
        "부서",
        "그룹",
    }
)


def _clean_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    return str(value).strip()


def _unique_preserve_order(values: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    unique_values: list[str] = []
    for value in values:
        candidate = value.strip()
        if not candidate or candidate in seen:
            continue
        seen.add(candidate)
        unique_values.append(candidate)
    return unique_values


def _update_group_targets(group_name: str, mutate: Callable[[list[str]], list[str]]) -> None:
    cleaned_group_name = _clean_text(group_name)
    store = load_target_store()
    updated_groups: list[CustomGroup] = []
    matched = False

    for group in store.custom_groups:
        if group.name != cleaned_group_name:
            updated_groups.append(group)
            continue
        matched = True
        updated_groups.append(
            CustomGroup(
                name=group.name,
                targets=_unique_preserve_order(mutate(list(group.targets))),
            )
        )

    if not matched:
        raise FolderScanError("선택한 그룹을 찾지 못했습니다.")

    updated_groups.sort(key=lambda group: group.name.casefold())
    save_custom_groups(updated_groups)


def add_custom_group(group_name: str) -> None:
    cleaned_group_name = _clean_text(group_name)
    if not cleaned_group_name:
        raise FolderScanError("그룹 이름을 입력해 주세요.")

    store = load_target_store()
    if cleaned_group_name in get_custom_group_names(store):
        raise FolderScanError("이미 같은 이름의 그룹이 있습니다.")

    updated_groups = [*store.custom_groups, CustomGroup(name=cleaned_group_name, targets=[])]
    updated_groups.sort(key=lambda group: group.name.casefold())
    save_custom_groups(updated_groups)


def rename_custom_group(old_group_name: str, new_group_name: str) -> None:
    cleaned_old_group_name = _clean_text(old_group_name)
    cleaned_new_group_name = _clean_text(new_group_name)
    if not cleaned_old_group_name:
        raise FolderScanError("변경할 기존 그룹을 선택해 주세요.")
    if not cleaned_new_group_name:
        raise FolderScanError("새 그룹 이름을 입력해 주세요.")
    if cleaned_old_group_name == cleaned_new_group_name:
        raise FolderScanError("새 이름이 현재 이름과 같습니다.")

    store = load_target_store()
    if cleaned_new_group_name in get_custom_group_names(store):
        raise FolderScanError("이미 같은 이름의 그룹이 있습니다.")

    updated_groups: list[CustomGroup] = []
    matched = False
    for group in store.custom_groups:
        if group.name != cleaned_old_group_name:
            updated_groups.append(group)
            continue
        matched = True
        updated_groups.append(CustomGroup(name=cleaned_new_group_name, targets=list(group.targets)))

    if not matched:
        raise FolderScanError("변경할 그룹을 찾지 못했습니다.")

    updated_groups.sort(key=lambda group: group.name.casefold())
    save_custom_groups(updated_groups)


def delete_custom_group(group_name: str) -> None:
    cleaned_group_name = _clean_text(group_name)
    store = load_target_store()
    updated_groups = [group for group in store.custom_groups if group.name != cleaned_group_name]
    if len(updated_groups) == len(store.custom_groups):
        raise FolderScanError("삭제할 그룹을 찾지 못했습니다.")
    save_custom_groups(updated_groups)


def add_target_to_custom_group(group_name: str, target_name: str) -> None:
    cleaned_target_name = _clean_text(target_name)
    if not cleaned_target_name:
        raise FolderScanError("추가할 대상명을 입력해 주세요.")

    def mutate(targets: list[str]) -> list[str]:
        if cleaned_target_name not in targets:
            targets.append(cleaned_target_name)
        return targets

    _update_group_targets(group_name, mutate)


def add_targets_to_custom_group(group_name: str, target_names: Sequence[str]) -> int:
    cleaned_targets = _unique_preserve_order(_clean_text(target) for target in target_names if _clean_text(target))
    if not cleaned_targets:
        raise FolderScanError("엑셀에서 읽을 수 있는 대상명이 없습니다.")

    added_count = 0

    def mutate(targets: list[str]) -> list[str]:
        nonlocal added_count
        for target_name in cleaned_targets:
            if target_name not in targets:
                targets.append(target_name)
                added_count += 1
        return targets

    _update_group_targets(group_name, mutate)
    return added_count


def delete_target_from_custom_group(group_name: str, target_name: str) -> None:
    cleaned_target_name = _clean_text(target_name)
    if not cleaned_target_name:
        raise FolderScanError("삭제할 대상명이 비어 있습니다.")

    removed = False

    def mutate(targets: list[str]) -> list[str]:
        nonlocal removed
        filtered_targets = [target for target in targets if target != cleaned_target_name]
        removed = len(filtered_targets) != len(targets)
        return filtered_targets

    _update_group_targets(group_name, mutate)
    if not removed:
        raise FolderScanError("삭제할 대상을 찾지 못했습니다.")


def delete_targets_from_custom_group(group_name: str, target_names: Sequence[str]) -> int:
    cleaned_targets = _unique_preserve_order(_clean_text(target) for target in target_names if _clean_text(target))
    if not cleaned_targets:
        raise FolderScanError("삭제할 대상을 하나 이상 선택해 주세요.")

    removed_count = 0

    def mutate(targets: list[str]) -> list[str]:
        nonlocal removed_count
        remaining_targets: list[str] = []
        removed_set = set(cleaned_targets)
        for target in targets:
            if target in removed_set:
                removed_count += 1
                continue
            remaining_targets.append(target)
        return remaining_targets

    _update_group_targets(group_name, mutate)
    return removed_count


def _looks_like_header(cell_value: object) -> bool:
    candidate = _clean_text(cell_value)
    if not candidate:
        return False
    normalized = candidate.replace(" ", "")
    if normalized in HEADER_HINTS:
        return True
    return any(hint in normalized for hint in ("대상", "이름", "명단", "소속", "부서"))


def import_targets_from_excel(file_storage: FileStorage, *, skip_header: bool) -> list[str]:
    filename = _clean_text(file_storage.filename)
    if not filename.lower().endswith(".xlsx"):
        raise FolderScanError("그룹 대상 업로드는 .xlsx 파일만 지원합니다.")

    stream = file_storage.stream
    if hasattr(stream, "seek"):
        stream.seek(0)

    try:
        workbook = load_workbook(stream, read_only=True, data_only=True)
    except Exception as exc:  # pragma: no cover - openpyxl error text varies
        raise FolderScanError("엑셀 파일을 열 수 없습니다. 파일 형식을 확인해 주세요.") from exc

    try:
        worksheet = workbook.worksheets[0]
        extracted_targets: list[str] = []

        for row_index, row in enumerate(worksheet.iter_rows(min_row=1, max_col=1, values_only=True), start=1):
            cell_value = row[0]
            if row_index == 1 and (skip_header or _looks_like_header(cell_value)):
                continue

            cleaned_value = _clean_text(cell_value)
            if cleaned_value and cleaned_value not in extracted_targets:
                extracted_targets.append(cleaned_value)

        if not extracted_targets:
            raise FolderScanError("엑셀 첫 번째 열에서 대상명을 찾지 못했습니다.")

        return extracted_targets
    finally:
        workbook.close()
